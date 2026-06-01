import zipfile, xml.etree.ElementTree as ET
import re
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

base = os.path.dirname(os.path.abspath(__file__))

# ── 1. Parse MAC tables from docx ──
def extract_mac_table(path):
    lines = []
    with zipfile.ZipFile(path, 'r') as z:
        with z.open('word/document.xml') as f:
            tree = ET.parse(f)
            ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
            for p in tree.iter('{%s}p' % ns):
                line = ''.join(t.text or '' for t in p.iter('{%s}t' % ns))
                if line.strip():
                    lines.append(line.strip())
    entries = []
    for l in lines:
        m = re.match(r'^([0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4})\s+(\d+)\s+Learned\s+(GE\d+/\d+/\d+)\s+([YN])', l, re.I)
        if m:
            entries.append({'mac': m.group(1).lower(), 'vlan': int(m.group(2)), 'port': m.group(3)})
    return entries

# ── 2. Parse VLAN config from xlsx ──
def extract_vlan_config(xlsx_path, sheet_idx, switch_name):
    """Extract port config from xlsx sheet. Returns dict: port -> {status, speed, pvid, vlan, type, desc, server}"""
    with zipfile.ZipFile(xlsx_path, 'r') as z:
        shared = []
        with z.open('xl/sharedStrings.xml') as f:
            tree = ET.parse(f)
            for si in tree.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}si'):
                t = si.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t')
                if t is not None:
                    shared.append(t.text or '')
                else:
                    texts = []
                    for r in si.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t'):
                        texts.append(r.text or '')
                    shared.append(''.join(texts))

        sheet_file = f'xl/worksheets/sheet{sheet_idx}.xml'
        with z.open(sheet_file) as f:
            tree = ET.parse(f)
            root = tree.getroot()
            ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
            rows = []
            for row in root.iter('{%s}row' % ns):
                cells = []
                for c in row.iter('{%s}c' % ns):
                    v = c.find('{%s}v' % ns)
                    if v is not None and v.text:
                        vt = c.get('t', '')
                        if vt == 's':
                            idx = int(v.text)
                            cells.append(shared[idx] if idx < len(shared) else '')
                        else:
                            cells.append(v.text)
                    else:
                        cells.append('')
                rows.append(cells)
    # Data rows: A=port, B=status, C=speed, D=duplex, E=type, F=pvid, G=vlan, H=desc, I=server, J=notes
    port_configs = {}
    for row in rows:
        if len(row) >= 1 and row[0] and re.match(r'^GE\d+/\d+/\d+$', row[0]):
            port = row[0]
            port_configs[port] = {
                'status': row[1] if len(row) > 1 else '',
                'speed': row[2] if len(row) > 2 else '',
                'duplex': row[3] if len(row) > 3 else '',
                'type': row[4] if len(row) > 4 else '',
                'pvid': row[5] if len(row) > 5 else '',
                'vlan': row[6] if len(row) > 6 else '',
                'desc': row[7] if len(row) > 7 else '',
                'server': row[8] if len(row) > 8 else '',
                'notes': row[9] if len(row) > 9 else '',
            }
    return port_configs

# ── 3. Load data ──
c06_mac = extract_mac_table(os.path.join(base, 'C06-H3C.docx'))
d03_mac = extract_mac_table(os.path.join(base, 'D03-H3C.docx'))
d04_mac = extract_mac_table(os.path.join(base, 'D04-H3C.docx'))

xlsx_path = os.path.join(base, '..', 'H3C-VLAN配置整理_v3.xlsx')
c06_vlan = extract_vlan_config(xlsx_path, 2, 'C06')  # Sheet 2 = C06
d03_vlan = extract_vlan_config(xlsx_path, 3, 'D03')  # Sheet 3 = D03
d04_vlan = extract_vlan_config(xlsx_path, 4, 'D04')  # Sheet 4 = D04

trunk_ports = {
    'C06': {'GE1/0/31': 'D04', 'GE1/0/48': 'USG6305E(极云)'},
    'D03': {'GE1/0/48': 'C06(上联)'},
    'D04': {'GE1/0/31': 'C06', 'GE1/0/48': 'D03'},
}

# ── 4. Build port-centric data ──
def build_port_data(switch_name, mac_entries, vlan_config, trunk):
    # Group MACs by port
    port_macs = {}
    for e in mac_entries:
        port = e['port']
        if port not in port_macs:
            port_macs[port] = []
        port_macs[port].append(e['mac'])

    # Build full port list (GE1/0/1 ~ GE1/0/52)
    result = []
    for i in range(1, 53):
        port = f'GE1/0/{i}'
        macs = sorted(set(port_macs.get(port, [])))
        vc = vlan_config.get(port, {})
        is_trunk = port in trunk
        result.append({
            'switch': switch_name,
            'port': port,
            'status': vc.get('status', ''),
            'speed': vc.get('speed', ''),
            'type': vc.get('type', ''),
            'pvid': vc.get('pvid', ''),
            'vlan': vc.get('vlan', ''),
            'desc': vc.get('desc', ''),
            'server': vc.get('server', ''),
            'notes': vc.get('notes', ''),
            'mac_count': len(macs),
            'macs': '\n'.join(macs) if macs else '',
            'is_trunk': is_trunk,
            'trunk_to': trunk.get(port, ''),
            'nic_name': '',      # user fills in
        })
    return result

c06_data = build_port_data('C06', c06_mac, c06_vlan, trunk_ports['C06'])
d03_data = build_port_data('D03', d03_mac, d03_vlan, trunk_ports['D03'])
d04_data = build_port_data('D04', d04_mac, d04_vlan, trunk_ports['D04'])

all_data = c06_data + d03_data + d04_data

print(f"C06: {len(c06_mac)} MACs, D03: {len(d03_mac)} MACs, D04: {len(d04_mac)} MACs")
print(f"Total ports: {len(all_data)}")

# ── 5. Create Excel ──
wb = Workbook()

hdr_font = Font(bold=True, size=11, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='1e40af')
hdr_fill_c06 = PatternFill('solid', fgColor='1e40af')
hdr_fill_d03 = PatternFill('solid', fgColor='b45309')
hdr_fill_d04 = PatternFill('solid', fgColor='047857')
cell_font = Font(size=10)
mac_font = Font(size=9, name='Consolas', color='1e40af')
fill_trunk = PatternFill('solid', fgColor='eef2ff')
fill_up = PatternFill('solid', fgColor='f0fdf4')
fill_down = PatternFill('solid', fgColor='fafafa')
fill_user = PatternFill('solid', fgColor='fefce8')  # yellow for user to fill
thin_border = Border(
    left=Side(style='thin', color='cbd5e1'),
    right=Side(style='thin', color='cbd5e1'),
    top=Side(style='thin', color='cbd5e1'),
    bottom=Side(style='thin', color='cbd5e1')
)

COLUMNS = [
    ('交换机', 8),
    ('端口', 12),
    ('链路状态', 8),
    ('速率', 8),
    ('端口类型', 8),
    ('PVID', 6),
    ('所属VLAN', 14),
    ('用途/描述', 18),
    ('服务器名称', 16),
    ('备注', 16),
    ('MAC数量', 8),
    ('MAC地址', 30),
    ('Trunk对端', 16),
    ('网卡名称 (手动填写)', 18),
]

def write_sheet(wb, title, data, hdr_fill):
    ws = wb.create_sheet(title)
    for col_idx, (name, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}1'

    for i, d in enumerate(data, 2):
        vals = [
            d['switch'], d['port'], d['status'], d['speed'], d['type'],
            d['pvid'], d['vlan'], d['desc'], d['server'], d['notes'],
            d['mac_count'], d['macs'], d['trunk_to'], d['nic_name'],
        ]
        user_cols = (14,)  # nic_name column only
        for col_idx, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=col_idx, value=v)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=(col_idx == 12))

            if col_idx == 12:  # MAC column
                cell.font = mac_font
            elif col_idx in user_cols:
                cell.fill = fill_user
                cell.font = Font(size=10, color='b45309')
            else:
                cell.font = cell_font

        # Row color by status
        row_fill = None
        if d['is_trunk']:
            row_fill = fill_trunk
        elif d['status'] == 'UP':
            row_fill = fill_up
        elif d['status'] == 'DOWN' or d['status'] == '':
            row_fill = fill_down

        if row_fill:
            for col_idx in range(1, len(COLUMNS) + 1):
                c = ws.cell(row=i, column=col_idx)
                if col_idx not in user_cols:
                    c.fill = row_fill

    # Set row height for MAC display
    for i, d in enumerate(data, 2):
        if d['mac_count'] > 1:
            ws.row_dimensions[i].height = max(15, d['mac_count'] * 14)

# ── Sheet 1: All switches combined ──
ws_all = wb.active
ws_all.title = 'All Ports'
for col_idx, (name, width) in enumerate(COLUMNS, 1):
    cell = ws_all.cell(row=1, column=col_idx, value=name)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border
    ws_all.column_dimensions[get_column_letter(col_idx)].width = width
ws_all.freeze_panes = 'A2'
ws_all.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}1'

for i, d in enumerate(all_data, 2):
    vals = [
        d['switch'], d['port'], d['status'], d['speed'], d['type'],
        d['pvid'], d['vlan'], d['desc'], d['server'], d['notes'],
        d['mac_count'], d['macs'], d['trunk_to'], d['nic_name'],
    ]
    user_cols = (14,)
    for col_idx, v in enumerate(vals, 1):
        cell = ws_all.cell(row=i, column=col_idx, value=v)
        cell.border = thin_border
        cell.alignment = Alignment(vertical='center', wrap_text=(col_idx == 12))
        if col_idx == 12:
            cell.font = mac_font
        elif col_idx in user_cols:
            cell.fill = fill_user
            cell.font = Font(size=10, color='b45309')
        else:
            cell.font = cell_font

    row_fill = None
    if d['is_trunk']:
        row_fill = fill_trunk
    elif d['status'] == 'UP':
        row_fill = fill_up
    else:
        row_fill = fill_down
    if row_fill:
        for col_idx in range(1, len(COLUMNS) + 1):
            c = ws_all.cell(row=i, column=col_idx)
            if col_idx not in user_cols:
                c.fill = row_fill

    if d['mac_count'] > 1:
        ws_all.row_dimensions[i].height = max(15, d['mac_count'] * 14)

# ── Sheet 2-4: Per switch ──
write_sheet(wb, 'C06 Ports', c06_data, hdr_fill_c06)
write_sheet(wb, 'D03 Ports', d03_data, hdr_fill_d03)
write_sheet(wb, 'D04 Ports', d04_data, hdr_fill_d04)

# ── Sheet 5: Summary - only ports with MACs ──
ws_sum = wb.create_sheet('有MAC的端口')
sum_cols = [
    ('交换机', 8), ('端口', 12), ('链路状态', 8), ('端口类型', 8),
    ('VLAN', 14), ('服务器名称', 16), ('MAC数量', 8), ('MAC地址', 30),
    ('网卡名称 (手动填写)', 18),
]
for col_idx, (name, width) in enumerate(sum_cols, 1):
    cell = ws_sum.cell(row=1, column=col_idx, value=name)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border
    ws_sum.column_dimensions[get_column_letter(col_idx)].width = width
ws_sum.freeze_panes = 'A2'

filtered = [d for d in all_data if d['mac_count'] > 0]
for i, d in enumerate(filtered, 2):
    vals = [d['switch'], d['port'], d['status'], d['type'], d['vlan'],
            d['server'], d['mac_count'], d['macs'], d['nic_name']]
    user_cols = (9,)
    for col_idx, v in enumerate(vals, 1):
        cell = ws_sum.cell(row=i, column=col_idx, value=v)
        cell.border = thin_border
        cell.alignment = Alignment(vertical='center', wrap_text=(col_idx == 8))
        if col_idx == 8:
            cell.font = mac_font
        elif col_idx in user_cols:
            cell.fill = fill_user
            cell.font = Font(size=10, color='b45309')
        else:
            cell.font = cell_font
    row_fill = fill_trunk if d['is_trunk'] else (fill_up if d['status'] == 'UP' else fill_down)
    for col_idx in range(1, len(sum_cols) + 1):
        c = ws_sum.cell(row=i, column=col_idx)
        if col_idx not in user_cols:
            c.fill = row_fill
    if d['mac_count'] > 1:
        ws_sum.row_dimensions[i].height = max(15, d['mac_count'] * 14)

output = os.path.join(base, 'H3C-端口MAC台账.xlsx')
wb.save(output)
print(f'\nSaved: {output}')
print(f'Ports with MAC: {len(filtered)} / {len(all_data)}')
