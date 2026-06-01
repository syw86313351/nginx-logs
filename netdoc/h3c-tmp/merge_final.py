import zipfile, xml.etree.ElementTree as ET
import re, os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

base = os.path.dirname(os.path.abspath(__file__))

# ── 1. Parse MAC tables ──
def extract_mac_table(path):
    lines = []
    with zipfile.ZipFile(path, 'r') as z:
        with z.open('word/document.xml') as f:
            tree = ET.parse(f)
            ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
            for p in tree.iter('{%s}p' % ns):
                line = ''.join(t.text or '' for t in p.iter('{%s}t' % ns))
                if line.strip():
                    lines.append(line.strip())
    entries = []
    for l in lines:
        m = re.match(r'^([0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4})\s+(\d+)\s+Learned\s+(GE\d+/\d+/\d+)\s+([YN])', l, re.I)
        if m:
            entries.append({'mac': m.group(1).lower(), 'vlan': int(m.group(2)), 'port': m.group(3)})
    return entries

# ── 2. Read original xlsx ──
def read_xlsx(path):
    with zipfile.ZipFile(path, 'r') as z:
        shared = []
        with z.open('xl/sharedStrings.xml') as f:
            tree = ET.parse(f)
            for si in tree.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}si'):
                t = si.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t')
                if t is not None:
                    shared.append(t.text or '')
                else:
                    texts = []
                    for r in si.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t'):
                        texts.append(r.text or '')
                    shared.append(''.join(texts))

        sheets_data = {}
        for i in range(1, 6):
            sheet_file = f'xl/worksheets/sheet{i}.xml'
            try:
                with z.open(sheet_file) as f:
                    tree = ET.parse(f)
                    root = tree.getroot()
                    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
                    rows = []
                    for row in root.iter('{%s}row' % ns):
                        cells = []
                        for c in row.iter('{%s}c' % ns):
                            v = c.find('{%s}v' % ns)
                            if v is not None and v.text:
                                vt = c.get('t', '')
                                if vt == 's':
                                    idx = int(v.text)
                                    cells.append(shared[idx] if idx < len(shared) else '')
                                else:
                                    cells.append(v.text)
                            else:
                                cells.append('')
                        rows.append(cells)
                    sheets_data[i] = rows
            except:
                pass
    return sheets_data

# ── 3. Load data ──
c06_mac = extract_mac_table(os.path.join(base, 'C06-H3C（mac）.docx'))
d03_mac = extract_mac_table(os.path.join(base, 'D03-H3C（mac）.docx'))
d04_mac = extract_mac_table(os.path.join(base, 'D04-H3C（mac）.docx'))

xlsx_path = os.path.join(base, '..', 'H3C-VLAN配置整理_v3.xlsx')
orig_data = read_xlsx(xlsx_path)

# Group MACs by switch+port
def group_macs(mac_entries):
    port_macs = {}
    for e in mac_entries:
        key = e['port']
        if key not in port_macs:
            port_macs[key] = []
        port_macs[key].append(e['mac'])
    return port_macs

c06_macs = group_macs(c06_mac)
d03_macs = group_macs(d03_mac)
d04_macs = group_macs(d04_mac)

trunk_ports = {
    'C06': {'GE1/0/31': 'D04', 'GE1/0/48': 'USG6305E(极云)'},
    'D03': {'GE1/0/48': 'C06(上联)'},
    'D04': {'GE1/0/31': 'C06', 'GE1/0/48': 'D03'},
}

print(f"C06: {len(c06_mac)} MACs, D03: {len(d03_mac)} MACs, D04: {len(d04_mac)} MACs")

# ── 4. Build new workbook ──
wb = Workbook()

# Styles
hdr_font = Font(bold=True, size=11, color='FFFFFF')
hdr_fill_s1 = PatternFill('solid', fgColor='1e3a5f')
hdr_fill_c06 = PatternFill('solid', fgColor='1e40af')
hdr_fill_d03 = PatternFill('solid', fgColor='b45309')
hdr_fill_d04 = PatternFill('solid', fgColor='047857')
hdr_fill_trunk = PatternFill('solid', fgColor='6d28d9')
hdr_fill_sum = PatternFill('solid', fgColor='b91c1c')
cell_font = Font(size=10)
mac_font = Font(size=9, name='Consolas', color='1e40af')
fill_trunk = PatternFill('solid', fgColor='eef2ff')
fill_up = PatternFill('solid', fgColor='f0fdf4')
fill_down = PatternFill('solid', fgColor='fafafa')
fill_user = PatternFill('solid', fgColor='fefce8')
thin_border = Border(
    left=Side(style='thin', color='cbd5e1'),
    right=Side(style='thin', color='cbd5e1'),
    top=Side(style='thin', color='cbd5e1'),
    bottom=Side(style='thin', color='cbd5e1')
)

def write_table(ws, start_row, rows, hdr_fill, mac_col_idx=None, user_col_idxs=None):
    """Write rows to worksheet with formatting."""
    if not rows:
        return start_row
    if user_col_idxs is None:
        user_col_idxs = []

    for r_idx, row in enumerate(rows):
        excel_row = start_row + r_idx
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=excel_row, column=c_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=(c_idx == mac_col_idx))

            if r_idx == 0:  # header row
                cell.font = hdr_font
                cell.fill = hdr_fill
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            elif c_idx == mac_col_idx:
                cell.font = mac_font
            elif c_idx in user_col_idxs:
                cell.fill = fill_user
                cell.font = Font(size=10, color='b45309')
            else:
                cell.font = cell_font

    return start_row + len(rows)

def write_switch_sheet(wb, title, orig_rows, mac_port_map, trunk, hdr_fill):
    """Write a switch sheet: original columns + MAC columns."""
    ws = wb.create_sheet(title)

    # Original header (row 1-3 from xlsx are metadata + header)
    # Row 1: title, Row 2: stats, Row 3: column headers
    # Data starts at row 4

    # Write original rows first
    for r_idx, row in enumerate(orig_rows):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx + 1, column=c_idx, value=val)
            cell.border = thin_border
            cell.font = cell_font
            cell.alignment = Alignment(vertical='center')

            if r_idx <= 1:  # title rows
                cell.font = Font(bold=True, size=11)
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif r_idx == 2:  # header row
                cell.font = hdr_font
                cell.fill = hdr_fill
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Add MAC columns to header row (row 3)
    orig_cols = len(orig_rows[2]) if len(orig_rows) > 2 else 8
    extra_headers = ['MAC数量', 'MAC地址', 'Trunk对端', '网卡名称(手动填写)']
    extra_widths = [8, 35, 16, 18]
    for i, (h, w) in enumerate(zip(extra_headers, extra_widths)):
        col = orig_cols + 1 + i
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = w

    # Add MAC data to each port row (starting at row 4)
    for r_idx in range(3, len(orig_rows)):
        row = orig_rows[r_idx]
        if not row or not row[0]:
            continue
        port = row[0]
        if not re.match(r'^GE\d+/\d+/\d+$', str(port)):
            continue

        macs = sorted(set(mac_port_map.get(port, [])))
        is_trunk = port in trunk
        mac_count = len(macs)
        mac_str = '\n'.join(macs) if macs else ''
        trunk_to = trunk.get(port, '') if is_trunk else ''

        excel_row = r_idx + 1
        # MAC count
        cell = ws.cell(row=excel_row, column=orig_cols + 1, value=mac_count)
        cell.border = thin_border
        cell.font = cell_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        # MAC list
        cell = ws.cell(row=excel_row, column=orig_cols + 2, value=mac_str)
        cell.border = thin_border
        cell.font = mac_font
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        # Trunk peer
        cell = ws.cell(row=excel_row, column=orig_cols + 3, value=trunk_to)
        cell.border = thin_border
        cell.font = cell_font
        cell.alignment = Alignment(vertical='center')
        # NIC name (user fills)
        cell = ws.cell(row=excel_row, column=orig_cols + 4)
        cell.border = thin_border
        cell.fill = fill_user
        cell.font = Font(size=10, color='b45309')
        cell.alignment = Alignment(vertical='center')

        # Row fill based on status
        status = row[1] if len(row) > 1 else ''
        if is_trunk:
            row_fill = fill_trunk
        elif status == 'UP':
            row_fill = fill_up
        else:
            row_fill = fill_down
        for c in range(1, orig_cols + 5):
            cc = ws.cell(row=excel_row, column=c)
            if c <= orig_cols:
                cc.fill = row_fill
            elif c == orig_cols + 4:
                pass  # keep user fill
            else:
                cc.fill = row_fill

        # Row height for multiple MACs
        if mac_count > 1:
            ws.row_dimensions[excel_row].height = max(15, mac_count * 14)

    # Set column widths for original columns
    default_widths = [12, 8, 8, 8, 8, 6, 14, 18]
    for i, w in enumerate(default_widths):
        if i < orig_cols:
            ws.column_dimensions[get_column_letter(i + 1)].width = w

    ws.freeze_panes = 'A4'

# ── Sheet 1: Overview (copy from original) ──
ws1 = wb.active
ws1.title = 'VLAN总览'
for r_idx, row in enumerate(orig_data.get(1, [])):
    for c_idx, val in enumerate(row, 1):
        cell = ws1.cell(row=r_idx + 1, column=c_idx, value=val)
        cell.border = thin_border
        cell.font = cell_font
        cell.alignment = Alignment(vertical='center')
        if r_idx <= 1:
            cell.font = Font(bold=True, size=11)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif r_idx == 2:
            cell.font = hdr_font
            cell.fill = hdr_fill_s1
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws1.freeze_panes = 'A4'

# ── Sheet 2-4: Switch details with MAC ──
write_switch_sheet(wb, 'C06 端口详情+MAC', orig_data.get(2, []), c06_macs, trunk_ports['C06'], hdr_fill_c06)
write_switch_sheet(wb, 'D03 端口详情+MAC', orig_data.get(3, []), d03_macs, trunk_ports['D03'], hdr_fill_d03)
write_switch_sheet(wb, 'D04 端口详情+MAC', orig_data.get(4, []), d04_macs, trunk_ports['D04'], hdr_fill_d04)

# ── Sheet 5: VLAN stats (copy from original) ──
ws5 = wb.create_sheet('VLAN端口统计')
for r_idx, row in enumerate(orig_data.get(5, [])):
    for c_idx, val in enumerate(row, 1):
        cell = ws5.cell(row=r_idx + 1, column=c_idx, value=val)
        cell.border = thin_border
        cell.font = cell_font
        cell.alignment = Alignment(vertical='center')
        if r_idx <= 1:
            cell.font = Font(bold=True, size=11)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif r_idx == 2:
            cell.font = hdr_font
            cell.fill = hdr_fill_s1
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws5.freeze_panes = 'A4'

# ── Sheet 6: Trunk links (copy from original) ──
if 1 in orig_data:
    # Find trunk section in sheet 1
    trunk_rows = []
    in_trunk = False
    for row in orig_data[1]:
        if row and 'Trunk' in str(row[0]):
            in_trunk = True
        if in_trunk:
            trunk_rows.append(row)

    if trunk_rows:
        ws6 = wb.create_sheet('Trunk互联链路')
        for r_idx, row in enumerate(trunk_rows):
            for c_idx, val in enumerate(row, 1):
                cell = ws6.cell(row=r_idx + 1, column=c_idx, value=val)
                cell.border = thin_border
                cell.font = cell_font
                cell.alignment = Alignment(vertical='center')
                if r_idx <= 1:
                    cell.font = hdr_font
                    cell.fill = hdr_fill_trunk
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws6.freeze_panes = 'A3'

# ── Sheet 7: Summary - ports with MAC ──
ws_sum = wb.create_sheet('有MAC的端口(汇总)')
sum_headers = ['交换机', '端口', '链路状态', '端口类型', 'PVID', 'VLAN', '服务器名称', 'MAC数量', 'MAC地址', 'Trunk对端', '网卡名称(手动填写)']
sum_widths = [8, 12, 8, 8, 6, 14, 16, 8, 35, 16, 18]

for c_idx, (h, w) in enumerate(zip(sum_headers, sum_widths), 1):
    cell = ws_sum.cell(row=1, column=c_idx, value=h)
    cell.font = hdr_font
    cell.fill = hdr_fill_sum
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border
    ws_sum.column_dimensions[get_column_letter(c_idx)].width = w
ws_sum.freeze_panes = 'A2'

# Build summary data from all three switches
all_switches = [
    ('C06', orig_data.get(2, []), c06_macs, trunk_ports['C06']),
    ('D03', orig_data.get(3, []), d03_macs, trunk_ports['D03']),
    ('D04', orig_data.get(4, []), d04_macs, trunk_ports['D04']),
]

row_num = 2
for switch_name, orig_rows, mac_map, trunk in all_switches:
    for r_idx in range(3, len(orig_rows)):
        row = orig_rows[r_idx]
        if not row or not row[0]:
            continue
        port = row[0]
        if not re.match(r'^GE\d+/\d+/\d+$', str(port)):
            continue

        macs = sorted(set(mac_map.get(port, [])))
        if not macs:
            continue

        status = row[1] if len(row) > 1 else ''
        ptype = row[4] if len(row) > 4 else ''
        pvid = row[5] if len(row) > 5 else ''
        vlan = row[6] if len(row) > 6 else ''
        server = row[8] if len(row) > 8 else ''
        is_trunk = port in trunk
        trunk_to = trunk.get(port, '') if is_trunk else ''

        vals = [switch_name, port, status, ptype, pvid, vlan, server,
                len(macs), '\n'.join(macs), trunk_to, '']

        for c_idx, v in enumerate(vals, 1):
            cell = ws_sum.cell(row=row_num, column=c_idx, value=v)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=(c_idx == 9))

            if c_idx == 9:
                cell.font = mac_font
            elif c_idx == 11:
                cell.fill = fill_user
                cell.font = Font(size=10, color='b45309')
            else:
                cell.font = cell_font

        # Row fill
        if is_trunk:
            rfill = fill_trunk
        elif status == 'UP':
            rfill = fill_up
        else:
            rfill = fill_down
        for c_idx in range(1, 12):
            cc = ws_sum.cell(row=row_num, column=c_idx)
            if c_idx != 11:
                cc.fill = rfill

        if len(macs) > 1:
            ws_sum.row_dimensions[row_num].height = max(15, len(macs) * 14)

        row_num += 1

# Save
output = os.path.join(base, '..', 'H3C-VLAN配置整理_v3-最终版.xlsx')
wb.save(output)
print(f'Saved: {output}')
print(f'Summary rows: {row_num - 2}')
